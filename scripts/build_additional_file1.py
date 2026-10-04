import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from pathlib import Path
import argparse
parser=argparse.ArgumentParser(description="Build the revision-2 supplementary workbook from the archived CSVs (see results/revision2/README.md for audit status).")
parser.add_argument('--results', type=Path, default=Path(__file__).resolve().parents[1]/'results/revision2')
parser.add_argument('--out', type=Path, default=Path('Additional_file_1.xlsx'))
args=parser.parse_args()
if args.out.exists(): parser.error('Output already exists; choose a new --out path')
U=args.results
supp=U/'supplementary_tables_S0-S5'
S0=pd.read_csv(supp/'S0_per_run.csv'); S0=S0[S0.feature_method!='Variance']
S0b=pd.read_csv(supp/'S0b_rev2_per_run.csv')
Sm=[pd.read_csv(supp/f) for f in ['S1_ARI_per_run.csv','S2_FM_per_run.csv','S3_Jaccard_per_run.csv','S4_Purity_per_run.csv']]
Sm=[x[x.feature_method!='Variance'] for x in Sm]
S5=pd.read_csv(supp/'S5_missing_configurations.csv'); S5=S5[S5.feature_method!='Variance']
S5b=pd.read_csv(supp/'S5b_rev2_missing.csv')
ov=pd.read_csv(U/'05_feature_sets_gate_ranking/feature_set_overlap_v2.csv'); ov=ov[~ov.method.isin(['Variance','DIFS_mc'])]
rt=pd.read_csv(U/'06_runtime/runtime_scaling.csv'); rt=rt[rt.fmethod.isin(['DIFS','Seurat','FEAST','GateOnly'])].copy()
rt.loc[rt['data.name']!='Baron',['maxrss_GB','mem_tasks']]=float('nan')
s8a=pd.read_csv(U/'07_simulation/E8_assumption.csv')
s8b=pd.read_csv(U/'07_simulation/sim_complete_per_panel.csv')
s9=pd.read_csv(U/'08_markers/marker_coverage_per_dataset.csv')
s9b=pd.read_csv(U/'08_markers/marker_coverage_tests.csv')
assert len(s8a)==300 and len(s8b)==1263 and len(s9)==3120 and len(s9b)==768, (len(s8a),len(s8b),len(s9),len(s9b))
assert len(S0)==2115, len(S0); assert len(S0b)==2097
assert S0[S0.result_set=='ncurve'].shape[0]==1852, S0[S0.result_set=='ncurve'].shape
for x in Sm: assert len(x)==1852, len(x)
readme=[
 ("Audit status", "2026-10-04 corrected source snapshot. Signed-rank differences are rounded to 12 decimal places before zero removal and midranking. This material update is not itself a tagged v2.1 release."),
 ("Additional file 1","Supplementary tables for: DIFS: Discriminative Feature Selection for Cell Clustering Based on Single-Cell RNA Sequencing Data"),
 ("",""),
 ("Sheet","Contents"),
 ("S0","Every run of the main grid (result_set = ncurve), the component analysis (abl), the reference-distribution comparison (hart) and the repaired Baron runs (baronfix). Filter on is_analysis_row = TRUE before summarising: the baronfix rows are provenance for the repaired Baron values and would otherwise count Baron twice. ARI_before_repair holds the Baron values before the SC3 wrapper was repaired."),
 ("S0b","Every run added in this revision: extended component analysis (abl2), repeated runs over clustering seeds and stratified 90% subsamples (reps), and a range of cluster numbers k_true + k_offset (krange). status = skipped_k_below_2 marks offsets that would give fewer than two clusters, which were not run."),
 ("S1-S4","Main grid by metric (ARI, Fowlkes-Mallows, Jaccard index on cell pairs, purity): dataset x method x clustering method x k policy x feature budget, with the realised feature count and a flag for repaired Baron values."),
 ("S5","Main-grid configurations that did not complete (18 FEAST on SimKumar4hard, 2 Seurat on Zhengmix8eq)."),
 ("S5b","Second-round tasks that did not complete (all in LAPACK dgesdd, all in competitor arms)."),
 ("S6","Cross-dataset agreement of the 1000-gene selections (estimated-k policy): both selections restricted to the genes measured in both datasets; exact hypergeometric chance baseline; jaccard_adj = chance-corrected Jaccard. Pairs with different gene identifier spaces are listed as not comparable."),
 ("S7","Median end-to-end wall time (minutes, true k, over the eight budgets) and median peak resident memory (GB, scheduler MaxRSS) per dataset, method and final clustering method. Memory is given for Baron only, from the repaired runs; for the short runs on the smaller datasets the scheduler recorded no usable peak, so those cells are left empty."),
 ("S8a","Simulation E8 (Appendix B): stage I alone against Seurat's standardised variance on panels whose expressed values are Gaussian within a cell type; one row per panel and ranking (60 panels: effect = delta/sigma in 2,4,6,8; n_cells 300/600/1200; 5 replicates). auc_vs_hvnull: AUC of markers against the unimodal genes of matched variance; top500_hv: those genes among the first 500."),
 ("S8b","Complete method against Seurat and FEAST on simulated panels (Appendix B, Table B4): one row per panel, method and budget (85 panels: E8, E8md, E4nd, E4; 600 cells; k = 5). recall / precision of the selected set for the markers; hv_selected: unimodal genes of matched variance selected (E8, E8md); from_stage2_only: markers selected by DIFS that were not in its own stage I list. GateOnly rows are expected values. FEAST has no rows on the four E4nd panels where all three seeds failed."),
 ("S9","Marker recovery on the 13 real datasets (Appendix C): one row per dataset, reference list (hurdle = primary; level and detection = its two parts; wilcox = comparison), method and budget. cov10 / cov25 / covall: share of the top 10 / top 25 per type / all qualifying markers contained in the selection; types1 / types3: share of annotated types with at least 1 / 3 of their top-10 markers selected; precision: share of selected genes that are markers of some type. Rows with an empty budget give need50 / need80, the number of top-ranked genes needed to cover 50% / 80% of the top-10 list, for methods with a full ranking. GateOnly and Random rows are expected values."),
 ("S9b","Paired tests for Appendix C: DIFS minus each comparator, per list, budget and metric, on the nine independent datasets and on all 13; exact sign-permutation signed-rank p-values (differences rounded to 12 decimals, zeros dropped, midranks), not adjusted for multiplicity. primary = TRUE marks the pre-specified comparison."),
 ("",""),
 ("Methods","DIFS: the method. Seurat: Seurat's variance-stabilising HVG selection. FEAST: FEAST. GateOnly: random selection from the genes passing the DIFS expression gate (internal control). DIFS_stage1only: DIFS with stage II disabled. DIFS_hartigan: DIFS with stage I calibrated by the original dip test. HVG_*: stage I ordered by standardised variance over the same gated genes. DIFS_fixed / DIFS_fixratio / DIFS_stage2only: DIFS with the MSE search removed / with the ratio search only / with stage II genes alone."),
]
out=args.out
out.parent.mkdir(parents=True, exist_ok=True)
with pd.ExcelWriter(out, engine='openpyxl') as w:
    pd.DataFrame(readme).to_excel(w,sheet_name='README',index=False,header=False)
    for name,df in [('S0',S0),('S0b',S0b),('S1_ARI',Sm[0]),('S2_FM',Sm[1]),('S3_Jaccard',Sm[2]),('S4_Purity',Sm[3]),('S5',S5),('S5b',S5b),('S6_overlap',ov),('S7_runtime_memory',rt),('S8a_sim_E8',s8a),('S8b_sim_complete',s8b),('S9_markers',s9),('S9b_marker_tests',s9b)]:
        df.to_excel(w,sheet_name=name,index=False)
wb=load_workbook(out)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row: c.font=Font(name='Arial',size=10,bold=c.font.bold)
    if ws.title!='README':
        for c in ws[1]: c.font=Font(name='Arial',size=10,bold=True); c.fill=PatternFill('solid',fgColor='E8EBE9')
        ws.freeze_panes='A2'
        for col in ws.columns:
            L=max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200]); ws.column_dimensions[col[0].column_letter].width=min(max(10,L+2),40)
    else:
        ws.column_dimensions['A'].width=18; ws.column_dimensions['B'].width=120
        for r in (1,3):
            for c in ws[r]: c.font=Font(name='Arial',size=11 if r==1 else 10,bold=True)
        for row in ws.iter_rows():
            for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
wb.save(out)
print({ws.title:ws.max_row-1 for ws in wb.worksheets})
